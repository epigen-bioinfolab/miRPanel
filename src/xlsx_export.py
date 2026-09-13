"""Spreadsheet exports.

The TSV files are built for R and pandas: one row per stored record,
provenance as '#' lines they skip, and full float precision on every number.
That is right for a script and wrong for a person - opened in a spreadsheet
the results file was 33 columns of alignment strings, and the enrichment file
was 232 rows of p-values printed to seventeen significant figures.

So these are a second format for a different reader, never a replacement.
Nothing here recomputes anything: if one of these files and its TSV ever
disagree, this module is wrong.

Results workbook
  Results     one row per gene, one column per tool
  All sites   every stored site, without the sequence strings
  Provenance  the same lines the TSV carries
  Tools run   what ran, how long, and what it said

Enrichment workbook
  one sheet per ontology, because each is tested against a different
  annotated background and must never share a ranking
  Provenance  the same lines the TSV carries
"""
import io

import report

REGION_LABEL = {"5UTR": "5′UTR", "CDS": "CDS", "3UTR": "3′UTR"}

SCI = "0.00E+00"

# One row per gene. Each tool keeps its own columns and its own scale;
# nothing is combined, and a tool that said nothing leaves an empty cell.
COLUMNS = [
    ("miRNA", 22), ("Gene", 12), ("Entrez", 10), ("RefSeq", 18),
    ("Regions hit", 16), ("Seed site", 12), ("Seed position", 16),
    ("Seed sites", 11), ("RNAhybrid ΔG", 15), ("ΔG position", 16),
    ("Duplexes", 10), ("TargetScan context++", 20), ("TargetScan PCT", 16),
    ("Conserved sites", 15), ("miRTarBase", 16), ("Strong functional", 17),
]

SITE_COLUMNS = [
    ("Gene", 12), ("RefSeq", 18), ("Region", 10), ("Tool", 13),
    ("Site type", 12), ("Start", 10), ("End", 10), ("Score", 12),
    ("Score is", 26), ("Low confidence", 15),
]

# k, n, K and N are the four integers of the hypergeometric test. The headers
# name them in words and keep the symbol, so a reader can follow the maths
# without the key and a statistician can still see which is which.
TERM_COLUMNS = [
    ("Term", 52), ("Term ID", 15),
    ("Genes found (k)", 15), ("Genes tested (n)", 16),
    ("In term, background (K)", 21), ("Background tested (N)", 20),
    ("p-value", 13, SCI), ("Adjusted p", 13, SCI), ("q-value", 13, SCI),
    ("Gene symbols", 42),
]

# Excel forbids ':' in a sheet name, and the panel spells these out anyway.
SHEET_NAME = {"GO:BP": "GO Biological Process",
              "GO:CC": "GO Cellular Component",
              "GO:MF": "GO Molecular Function",
              "KEGG": "KEGG"}


def _styles():
    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        raise RuntimeError(
            "openpyxl is not installed; run: pip install openpyxl")
    thin = Side(style="thin", color="D5DAE0")
    return {
        "head_font": Font(bold=True, size=10, color="FFFFFF"),
        "head_fill": PatternFill("solid", fgColor="3B4B5C"),
        # every cell centred both ways and wrapped, so a long label grows its
        # row instead of spilling across the neighbouring column
        "centre": Alignment(horizontal="center", vertical="center",
                            wrap_text=True),
        "left": Alignment(horizontal="left", vertical="top", wrap_text=True),
        "box": Border(left=thin, right=thin, top=thin, bottom=thin),
        "band": PatternFill("solid", fgColor="F7F9FA"),
        "Font": Font,
    }


def _sheet(ws, columns, data, s):
    """Write one table: styled header, centred cells, freeze and filter."""
    from openpyxl.utils import get_column_letter

    ws.append([c[0] for c in columns])
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = s["head_font"], s["head_fill"]
        cell.alignment, cell.border = s["centre"], s["box"]
        ws.column_dimensions[get_column_letter(i)].width = col[1]
    for row in data:
        ws.append(row)
    fmts = {i: c[2] for i, c in enumerate(columns, start=1) if len(c) > 2}
    for n, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for i, cell in enumerate(row, start=1):
            cell.alignment = s["centre"]
            cell.border = s["box"]
            if n % 2 == 0:
                cell.fill = s["band"]
            if i in fmts and isinstance(cell.value, (int, float)):
                cell.number_format = fmts[i]
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    if data:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(columns)),
                                          len(data) + 1)


def _provenance_sheet(wb, prov_lines, s):
    """The same lines the TSV carries, on their own sheet.

    Not above the data: in a spreadsheet, rows above the header are rows, and
    they break sorting, filtering and every formula that reads the column.
    """
    ps = wb.create_sheet("Provenance")
    ps.column_dimensions["A"].width = 26
    ps.column_dimensions["B"].width = 96
    for raw in prov_lines:
        text = raw[1:].strip() if raw.startswith("#") else raw
        key, sep, val = text.partition(":")
        # A provenance key is one word, because provenance.py writes it that
        # way. The prose notes also contain colons and must not be split into
        # a key nobody wrote.
        if sep and val.strip() and " " not in key.strip():
            ps.append([key.strip(), val.strip()])
        else:
            ps.append([None, text])
    for row in ps.iter_rows(min_col=1, max_col=2):
        row[0].font = s["Font"](bold=True, size=10)
        row[0].alignment = s["left"]
        row[1].alignment = s["left"]


# ------------------------------------------------------------------ results

def _gene_row(mirna, symbol, rows):
    """One gene's stored rows as one spreadsheet row."""
    g = report.summarise(rows)
    seed, rh, ts, mt = g["seed"], g["rh"], g["ts"], g["mt"]
    tx = (ts or {}).get("extra") or {}
    mx = (mt or {}).get("extra") or {}

    regions = " · ".join(REGION_LABEL.get(r, r) for r in g["regions"]) or None
    seed_pos = ("%s–%s %s" % (seed["pos_start"], seed["pos_end"],
                              REGION_LABEL.get(seed["region"], ""))
                if seed else None)
    rh_pos = ("%s–%s %s" % (rh["pos_start"], rh["pos_end"],
                            REGION_LABEL.get(rh["region"], ""))
              if rh else None)

    # TargetScan scores 3'UTRs. A gene with no 3'UTR evidence at all is
    # outside what it can judge, which is not the same as it looking and
    # finding nothing - so those cells say N.A. rather than sitting empty.
    if ts:
        context, pct, cons = (ts["score"], tx.get("aggregate_pct"),
                              tx.get("sites_conserved"))
    elif "3UTR" in g["regions"]:
        context = pct = cons = None
    else:
        context = pct = cons = "N.A."

    if mt:
        status = str(mx.get("status", "reported")).replace("_", " ")
        strong = mx.get("strong_functional")
    else:
        status = strong = None

    return [mirna, symbol, g["entrez"], g["refseq"], regions,
            seed["site_type"] if seed else None, seed_pos, g["n_seed"] or None,
            rh["score"] if rh else None, rh_pos, g["n_rh"] or None,
            context, pct, cons, status, strong]


def build(rows, prov_lines=(), tool_runs=(), title="Results"):
    """The results workbook, as bytes. `rows` are dicts with `extra` parsed."""
    from openpyxl import Workbook
    s = _styles()
    wb = Workbook()

    # group in the order the rows arrived, which is the order SQL sorted them
    by_gene, order = {}, []
    for r in rows:
        sym = r.get("symbol")
        if sym not in by_gene:
            by_gene[sym] = []
            order.append(sym)
        by_gene[sym].append(r)
    mirna = rows[0].get("mirna_name") if rows else None

    ws = wb.active
    ws.title = title
    _sheet(ws, COLUMNS, [_gene_row(mirna, g, by_gene[g]) for g in order], s)

    # Every site is still here, one per row - just without the four alignment
    # strings, which are a picture of a duplex rather than a value.
    _sheet(wb.create_sheet("All sites"), SITE_COLUMNS, [
        [r.get("symbol"), r.get("refseq"),
         REGION_LABEL.get(r.get("region"), r.get("region")), r.get("tool"),
         r.get("site_type"), r.get("pos_start"), r.get("pos_end"),
         r.get("score"), r.get("score_type"),
         "yes" if r.get("low_confidence") else ""] for r in rows], s)

    _provenance_sheet(wb, prov_lines, s)

    _sheet(wb.create_sheet("Tools run"),
           [("Tool", 16), ("Status", 11), ("Rows", 10), ("Seconds", 11),
            ("Message", 60)],
           [[t["tool"], t["status"], t["n_results"], t["seconds"],
             t["message"] or ""] for t in tool_runs], s)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------- enrichment

def build_enrichment(terms, prov_lines=()):
    """The enrichment workbook, as bytes. `terms` are enrichment.terms() rows.

    One sheet per ontology rather than one table with a source column. Each
    ontology is tested against a different annotated background - around
    17,116 genes for GO:BP but 8,516 for KEGG - so terms from two of them
    have never been comparable. Separate sheets make that structural instead
    of a warning nobody reads.
    """
    from openpyxl import Workbook
    s = _styles()
    wb = Workbook()
    wb.remove(wb.active)

    by_source, order = {}, []
    for t in terms:
        src = t["source"]
        if src not in by_source:
            by_source[src] = []
            order.append(src)
        by_source[src].append(t)

    for src in order:
        rows = [[t["description"], t["term_id"],
                 t["k_in_term_and_list"], t["n_in_list"],
                 t["k_in_term_and_background"], t["n_in_background"],
                 t["p_value"], t["p_adjust"], t["q_value"],
                 # stored slash-separated; commas read as a list to a person
                 (t["genes"] or "").replace("/", ", ")]
                for t in by_source[src]]
        name = SHEET_NAME.get(src) or str(src).replace(":", " ")
        _sheet(wb.create_sheet(name[:31]), TERM_COLUMNS, rows, s)

    if not order:
        _sheet(wb.create_sheet("No terms"), TERM_COLUMNS, [], s)

    _provenance_sheet(wb, prov_lines, s)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
